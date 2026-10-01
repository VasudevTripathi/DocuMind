from pathlib import Path
from typing import Dict, List, Any
import pypdf
import docx

class DocumentParsingError(Exception):
    """Raised when parsing a document fails."""
    pass

def parse_pdf(path: Path | str) -> Dict[str, Any]:
    """
    Extracts text from a PDF document using pypdf.
    Returns normalized dictionary with full text and page breakdown.
    """
    file_path = Path(path)
    if not file_path.exists():
        raise DocumentParsingError(f"PDF file not found: {file_path}")

    pages_data: List[Dict[str, Any]] = []
    full_text_parts: List[str] = []

    try:
        reader = pypdf.PdfReader(str(file_path))
        if len(reader.pages) == 0:
            raise DocumentParsingError("PDF document has 0 pages.")

        for idx, page in enumerate(reader.pages, start=1):
            page_text = page.extract_text() or ""
            pages_data.append({
                "page": idx,
                "text": page_text
            })
            if page_text.strip():
                full_text_parts.append(page_text)

        full_text = "\n\n".join(full_text_parts).strip()
        if not full_text:
            raise DocumentParsingError("No extractable text found in PDF (file may be scanned/image-only or encrypted).")

        return {
            "text": full_text,
            "pages": pages_data
        }
    except DocumentParsingError:
        raise
    except Exception as e:
        raise DocumentParsingError(f"Failed to parse PDF document: {str(e)}") from e

def parse_docx(path: Path | str) -> Dict[str, Any]:
    """
    Extracts text from a Word (.docx) document using python-docx.
    """
    file_path = Path(path)
    if not file_path.exists():
        raise DocumentParsingError(f"DOCX file not found: {file_path}")

    try:
        doc = docx.Document(str(file_path))
        text_blocks: List[str] = []

        # Extract paragraphs
        for para in doc.paragraphs:
            clean_para = para.text.strip()
            if clean_para:
                text_blocks.append(clean_para)

        # Extract table text
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    text_blocks.append(row_text)

        full_text = "\n\n".join(text_blocks).strip()
        if not full_text:
            raise DocumentParsingError("DOCX document contains no readable text.")

        return {
            "text": full_text,
            "pages": [{"page": 1, "text": full_text}]
        }
    except DocumentParsingError:
        raise
    except Exception as e:
        raise DocumentParsingError(f"Failed to parse DOCX document: {str(e)}") from e

def parse_txt(path: Path | str) -> Dict[str, Any]:
    """
    Extracts text from a plain text file (.txt).
    """
    file_path = Path(path)
    if not file_path.exists():
        raise DocumentParsingError(f"TXT file not found: {file_path}")

    for encoding in ["utf-8", "utf-8-sig", "latin-1", "cp1252"]:
        try:
            with open(file_path, "r", encoding=encoding) as f:
                content = f.read().strip()
                if not content:
                    raise DocumentParsingError("TXT file is empty.")
                return {
                    "text": content,
                    "pages": [{"page": 1, "text": content}]
                }
        except (UnicodeDecodeError, UnicodeError):
            continue
        except DocumentParsingError:
            raise
        except Exception as e:
            raise DocumentParsingError(f"Failed to read TXT file: {str(e)}") from e

    raise DocumentParsingError("Unable to decode text file with standard encodings.")

def parse_document(path: Path | str, file_type: str) -> Dict[str, Any]:
    """
    Dispatches document parsing based on file extension and detected type.
    """
    ext = Path(path).suffix.lower()
    if ext == ".pdf" or file_type.upper() == "PDF":
        return parse_pdf(path)
    elif ext in [".docx", ".doc"] or file_type.upper() == "DOCX":
        return parse_docx(path)
    elif ext == ".txt" or file_type.upper() == "TXT":
        return parse_txt(path)
    else:
        raise DocumentParsingError(f"Unsupported file type for automated text extraction: '{ext}'")
